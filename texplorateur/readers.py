import logging
import os
import warnings

import docx
import PyPDF2
import openpyxl
from PyPDF2.errors import PdfReadWarning

CONTEXTE_TAILLE = 60

# PyPDF2 émet des avertissements bruyants (logging + warnings.warn) pour des
# PDF malformés qu'il parvient quand même à lire (mode non strict) — ce ne
# sont pas des erreurs actionnables pour l'utilisateur d'une app GUI. Les
# émettre en masse depuis plusieurs threads en parallèle (recherche
# multithread sur des PDF) crée une forte contention — verrou interne du
# module logging, écritures stdout/stderr potentiellement bloquantes si le
# terminal ne suit pas — qui peut affamer le thread principal Tk au point
# de figer l'interface. On les désactive entièrement.
logging.getLogger("PyPDF2").setLevel(logging.ERROR)
warnings.filterwarnings("ignore", category=PdfReadWarning)


def lire_txt(chemin):
    try:
        with open(chemin, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    except Exception:
        return ""


def lire_docx(chemin):
    try:
        doc = docx.Document(chemin)
        return '\n'.join([p.text for p in doc.paragraphs])
    except Exception:
        return ""


def lire_pdf(chemin):
    try:
        texte = ""
        with open(chemin, 'rb') as f:
            lecteur = PyPDF2.PdfReader(f)
            for page in lecteur.pages:
                texte += page.extract_text() or ""
        return texte
    except Exception:
        return ""


def lire_xlsx(chemin):
    try:
        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        morceaux = []
        for feuille in classeur.worksheets:
            for ligne in feuille.iter_rows(values_only=True):
                for valeur in ligne:
                    if valeur is not None:
                        morceaux.append(str(valeur))
        return '\n'.join(morceaux)
    except Exception:
        return ""


LECTEURS_PAR_EXTENSION = {
    ".txt": lire_txt,
    ".docx": lire_docx,
    ".pdf": lire_pdf,
    ".xlsx": lire_xlsx,
}


def tronquer_chemin(chemin, largeur=70):
    """Raccourcit un chemin par le milieu pour éviter que l'affichage ne saute."""
    if len(chemin) <= largeur:
        return chemin
    garde = (largeur - 1) // 2
    return f"{chemin[:garde]}…{chemin[-garde:]}"


def extraire_contexte(contenu, phrase, taille=CONTEXTE_TAILLE):
    contenu_lower = contenu.lower()
    idx = contenu_lower.find(phrase.lower())
    if idx == -1:
        return ""
    debut = max(0, idx - taille)
    fin = min(len(contenu), idx + len(phrase) + taille)
    extrait = contenu[debut:fin].replace('\n', ' ').strip()
    prefixe = "…" if debut > 0 else ""
    suffixe = "…" if fin < len(contenu) else ""
    return f"{prefixe}{extrait}{suffixe}"
